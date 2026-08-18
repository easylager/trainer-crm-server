#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill Pavel's white-label DOCX template with Maksim Vasilenko content (no employer names)."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "Павел К.  Senior CV updated.docx"
OUT = ROOT / "cv-max-vasilenko-senior-python-whitelabel-ru-2026.docx"
OUT_NAMED = ROOT / "Василенко М. Senior CV.docx"


def set_runs_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def ensure_paragraphs(cell, count: int) -> None:
    while len(cell.paragraphs) < count:
        last = cell.paragraphs[-1]._p
        new_p = copy.deepcopy(last)
        for t in new_p.findall(".//" + qn("w:t")):
            t.text = ""
        last.addnext(new_p)
    while len(cell.paragraphs) > count:
        p = cell.paragraphs[-1]._p
        p.getparent().remove(p)


def _strip_num_pr(paragraph: Paragraph) -> None:
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return
    numPr = pPr.find(qn("w:numPr"))
    if numPr is not None:
        pPr.remove(numPr)


def _has_num_pr(paragraph: Paragraph) -> bool:
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return False
    return pPr.find(qn("w:numPr")) is not None


def cleanup_empty_bullets(cell) -> None:
    """Убрать пустые абзацы с маркером списка (квадратик/точка без текста)."""
    for paragraph in list(cell.paragraphs):
        if paragraph.text.strip():
            continue
        if not _has_num_pr(paragraph):
            continue
        parent = paragraph._p.getparent()
        # в ячейке должен остаться хотя бы один абзац
        if len(cell.paragraphs) > 1:
            parent.remove(paragraph._p)
        else:
            _strip_num_pr(paragraph)


def set_cell_lines(cell, lines: list[str]) -> None:
    ensure_paragraphs(cell, max(len(lines), 1))
    for i, line in enumerate(lines):
        set_runs_text(cell.paragraphs[i], line)
    # лишние пустые — удаляем целиком (иначе остаётся квадратный маркер списка)
    while len(cell.paragraphs) > len(lines):
        cell.paragraphs[-1]._p.getparent().remove(cell.paragraphs[-1]._p)
    cleanup_empty_bullets(cell)


def set_merged_value(row, col: int, text: str) -> None:
    cell = row.cells[col]
    if "\n" in text:
        set_cell_lines(cell, [x for x in text.split("\n") if x != ""])
    else:
        if not cell.paragraphs:
            cell.add_paragraph(text)
        else:
            set_runs_text(cell.paragraphs[0], text)
            for p in cell.paragraphs[1:]:
                set_runs_text(p, "")
        cleanup_empty_bullets(cell)


def set_space_after(paragraph: Paragraph, pt: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:after"), str(int(pt * 20)))  # twips
    # сохранить line из шаблона, если уже был
    if spacing.get(qn("w:line")) is None:
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")


def ensure_spacer_after_title(doc: Document) -> None:
    """Добавить воздух между 'Senior Python…' и верхней таблицей Навыки/Проекты."""
    title = doc.paragraphs[1]
    set_space_after(title, 12)
    title_p = title._p
    nxt = title_p.getnext()
    # если сразу таблица — вставить пустой абзац-спейсер
    if nxt is not None and nxt.tag == qn("w:tbl"):
        spacer = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "120")  # 6pt
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")
        pPr.append(spacing)
        spacer.append(pPr)
        title_p.addnext(spacer)


def _paragraph_text(p_elm) -> str:
    return "".join(t.text or "" for t in p_elm.findall(".//" + qn("w:t")))


def _make_spacer_paragraph(after_pt: float = 6, before_pt: float = 0, line: int = 240) -> OxmlElement:
    """Чистый пустой абзац без numPr/sectPr (иначе появляется □)."""
    spacer = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    if before_pt:
        spacing.set(qn("w:before"), str(int(before_pt * 20)))
    if after_pt:
        spacing.set(qn("w:after"), str(int(after_pt * 20)))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)
    spacer.append(pPr)
    return spacer


def remove_square_before_projects_heading(doc: Document) -> None:
    """Убрать □ над «Профессиональная деятельность (проекты)».

    В шаблоне перед заголовком пустой абзац с w:sectPr — в превью/Word
    он рисуется как пустой квадратик слева. Переносим sectPr в последний
    абзац резюме-таблицы, удаляем «грязные» пустые абзацы и ставим
    чистый спейсер, чтобы таблица проектов не прилипала к стеку.
    """
    body = doc.element.body
    children = list(body)
    heading_p = None
    heading_idx = None
    for i, el in enumerate(children):
        if el.tag != qn("w:p"):
            continue
        if _paragraph_text(el).strip() == "Профессиональная деятельность (проекты)":
            heading_p = el
            heading_idx = i
            break
    if heading_p is None or heading_idx is None:
        return

    sect_pr = None
    to_remove = []
    j = heading_idx - 1
    while j >= 0:
        el = children[j]
        if el.tag == qn("w:tbl"):
            break
        if el.tag == qn("w:p"):
            if _paragraph_text(el).strip():
                break
            found = el.find(".//" + qn("w:sectPr"))
            if found is not None and sect_pr is None:
                # отвязать от пустого абзаца
                parent = found.getparent()
                parent.remove(found)
                sect_pr = found
            to_remove.append(el)
        j -= 1

    if sect_pr is not None:
        # последний абзац ячейки summary (таблица перед пустыми)
        summary_cell = doc.tables[1].rows[0].cells[0]
        last_p = summary_cell.paragraphs[-1]._p
        pPr = last_p.get_or_add_pPr()
        # sectPr должен быть последним ребёнком pPr
        existing = pPr.find(qn("w:sectPr"))
        if existing is not None:
            pPr.remove(existing)
        pPr.append(sect_pr)

    for el in to_remove:
        body.remove(el)

    # воздух между стеком технологий и заголовком/таблицей проектов
    heading_p.addprevious(_make_spacer_paragraph(after_pt=6, before_pt=10, line=276))
    # чуть отступа и у самого заголовка
    h_pPr = heading_p.get_or_add_pPr()
    spacing = h_pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        h_pPr.insert(0, spacing)
    spacing.set(qn("w:before"), "120")  # 6pt
    spacing.set(qn("w:after"), "120")
    if spacing.get(qn("w:line")) is None:
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")


SKILLS = [
    "Навыки",
    "Python / Go API в production",
    "OAuth2, JWT, проверка прав на сервере",
    "Изоляция данных, least privilege, аудит",
    "FastAPI, Django/DRF, asyncio",
    "Kafka, gRPC, очереди сообщений",
    "PostgreSQL, Redis, оптимизация SQL",
    "AWS, GCP, Azure, Kubernetes",
    "pytest, CI/CD, code review",
]

PROJECT_NAMES = [
    "Проекты",
    "Realtime AI-ассистент для страховых кол-центров",
    "Международный SaaS: страховые заявки и antifraud",
    "Высоконагруженная страховая платформа (США)",
    "Платформа интеграций доставки",
    "Образовательная платформа · платежи и учёт",
]

SUMMARY = [
    "6+ лет коммерческой backend-разработки на Python: production API, микросервисы и интеграции, в том числе в регулируемых контурах с чувствительными данными. Практический опыт Golang на критичных участках распределённой платформы.",
    "API и сервисы: FastAPI, Django / DRF, Flask, asyncio, Pydantic, SQLAlchemy, Alembic, Celery, httpx, aiohttp, typing, OpenAPI/Swagger — проектирование и сопровождение контрактов под нагрузкой.",
    "Доступ и авторизация: OAuth2, валидация JWT, проверка прав на стороне сервера; при ошибке или неоднозначном результате проверки запрос отклоняется, без «пропуска по умолчанию». Мультитенантная изоляция данных и разграничение ролей.",
    "Надёжность границ API: идемпотентность, контроль входных данных, минимальные привилегии сервисных учёток, auditable-переходы статусов; на code review и в дизайне фич отдельно смотрю риски несанкционированного доступа и утечки данных между клиентами.",
    "Данные и обмен: PostgreSQL, MySQL, Redis; Kafka, gRPC, GCP Pub/Sub, Azure Service Bus, RabbitMQ; Airflow для пакетной обработки.",
    "Облака и поставка: AWS (EKS, IAM, RDS, S3, CloudWatch), GCP (Cloud Run, Pub/Sub, Cloud SQL), Azure (AKS, Blob, Service Bus); Docker, Kubernetes, ArgoCD, GitHub Actions / GitLab CI/CD; Grafana, Prometheus, Sentry.",
    "Недавний контур также включал AI-агентов и RAG в production (OpenAI / Claude, tool calling, guardrails, human-in-the-loop) — с жёсткими ограничениями на то, что автоматизация может сделать без подтверждения оператора.",
]

# Хронология подряд (новый → старый), без пересечений.
# AI ~19м | EU ~9м | US ~27м (самый длинный) | Delivery ~10м | Edu+Payments ~14м
PROJECTS = [
    {
        "from": "С 01.2025",
        "to": "До 07.2026",
        "role": "Инженер-программист",
        "project": (
            "Мультитенантный B2B SaaS для realtime AI-ассистенса в страховых кол-центрах. "
            "Во время звонка система ведёт оператора по сценарию: транскрибация, извлечение намерений, "
            "подсказки, вызов инструментов и ответы с опорой на базу знаний. Критичны низкая задержка, "
            "изоляция клиентов платформы и то, чтобы автоматизация не подставляла оператора на рискованных решениях."
        ),
        "duties": [
            "Отвечал за backend realtime-контура: приём аудио/событий звонка, статусы сессии, выдача подсказок оператору.",
            "Спроектировал мультитенантную архитектуру с изоляцией данных и прав доступа между клиентами платформы — запросы и retrieval не пересекают границы тенанта.",
            "В API ассистента проверял сессию/токен и роль оператора до выдачи подсказок и вызова инструментов; при сбое или неоднозначной проверке — отказ, без обхода «по умолчанию».",
            "Разрабатывал AI-агентские сценарии: многошаговые workflow, вызов инструментов, память сессии; на рискованных шагах — обязательное подтверждение человеком.",
            "Построил production RAG: ingest/chunking, embeddings, retrieval (pgvector), grounded-ответы LLM с опорой на источник и аудитом, что было показано оператору.",
            "Разделил sync API и тяжёлые AI-стадии через GCP Pub/Sub: ретраи, back-pressure, correlation ID и auditable-статусы шагов.",
            "Прорабатывал сценарии злоупотребления инструментами ассистента и утечки контекста между сессиями; оставлял ограничения и наблюдаемость на этих границах.",
            "Оптимизировал задержку подсказок: кеш в Redis, контроль стоимости LLM-вызовов, устойчивость при сбоях внешних API.",
            "Проектировал схему PostgreSQL под сессии, права, retrieval-контекст и аудит действий оператора.",
            "Настроил CI/CD, Docker/Kubernetes, Prometheus/Grafana/Sentry; покрыл ключевые сценарии pytest, включая негативные проверки доступа.",
        ],
        "tech": (
            "Python, Django, Django REST Framework, FastAPI, asyncio, PostgreSQL, pgvector, Redis, "
            "OAuth2/JWT, OpenAI API, Anthropic Claude, RAG, GCP (Cloud Run, Pub/Sub, Cloud SQL, GCS), "
            "Kubernetes, Docker, GitLab CI/CD, Prometheus, Grafana, Pytest, Sentry."
        ),
    },
    {
        "from": "С 04.2024",
        "to": "До 12.2024",
        "role": "Инженер-программист",
        "project": (
            "Международный B2B SaaS для обработки страховых заявок и выявления мошенничества (antifraud). "
            "Платформа принимает документы по заявке, извлекает данные, помогает проверить кейс и сигнализирует "
            "о подозрительных паттернах. Важны проверка человеком, партнёрские интеграции и надёжная доставка статусов."
        ),
        "duties": [
            "Разработал backend обработки страховых заявок: загрузка документов, извлечение полей, статусы для проверки оператором.",
            "Реализовал antifraud-сигналы и AI-подсказки по заявке с оценкой уверенности; финальное решение по рискованным кейсам оставалось за человеком.",
            "Построил event-driven потоки на Kafka для статусов заявок и webhook-уведомлений партнёрам.",
            "Интегрировал внешних партнёров: версионируемые контракты, подпись/проверка webhook, идемпотентная доставка статусов.",
            "Ограничивал доступ к вложениям и полям заявки по роли и принадлежности к организации клиента платформы.",
            "Участвовал в миграции Django → FastAPI на живом продукте без простоя; REST API с OpenAPI-спецификацией.",
            "Вынес тяжёлую обработку документов и antifraud-стадии в отдельные воркеры со служебными правами уже, чем у публичного API.",
            "Развивал схему PostgreSQL, оптимизировал SQL, кешировал горячие read-модели в Redis; покрыл сценарии pytest, в том числе отказы при невалидных credentials.",
        ],
        "tech": (
            "Python, FastAPI, Django, Django REST Framework, PostgreSQL, Redis, Kafka, Airflow, "
            "Azure (Service Bus, Blob Storage, AKS), Pydantic, Alembic, Docker, Kubernetes, Pytest, OpenAPI/Swagger, GitLab CI/CD."
        ),
    },
    {
        "from": "С 01.2022",
        "to": "До 03.2024",
        "role": "Инженер-программист",
        "project": (
            "Крупная высоконагруженная платформа в регулируемой среде (США): микросервисы для проверки "
            "доступности услуг и синхронизации клиентских данных. Система обслуживает миллионы пользователей, "
            "где ошибка в статусе сразу бьёт по клиенту. Критичны надёжность на сезонных пиках нагрузки, "
            "аудит изменений, быстрый разбор инцидентов и поставка в AWS/Kubernetes."
        ),
        "duties": [
            "Разрабатывал и поддерживал production API и мультисервисный пайплайн проверки статусов на Kafka — синхронизация данных между сервисами.",
            "Сопровождал плотный gRPC-трафик между соседними микросервисами; на стыке Python и Golang укреплял контракты и обработку ошибок авторизации вызовов.",
            "Участвовал в переносе отдельных стадий пайплайна на Python без ослабления проверок доступа между сервисами.",
            "Эксплуатировал микросервисы в AWS EKS/Kubernetes; работал с IAM и минимальными правами сервисных ролей; деплой через ArgoCD.",
            "Соблюдал практики безопасной работы с чувствительными клиентскими данными: least privilege, аудит изменений, наблюдаемость.",
            "На границах API закладывал отказ при невалидном или просроченном токене/контексте вызова — без «тихого» продолжения обработки.",
            "Оптимизировал горячие пути и согласование статусов; применял параллельную обработку под пиковую нагрузку.",
            "Поддерживал backend API и контракты данных для внутренних React-инструментов операционных команд с разграничением доступов.",
            "Писал автотесты (в т.ч. негативные) и собирал доказательства для приёмки; участвовал в разборе production-инцидентов.",
        ],
        "tech": (
            "Python, Golang, Flask, Kafka, gRPC, PostgreSQL, MySQL, SQLAlchemy, JWT/OAuth2, "
            "AWS (EKS, EC2, S3, RDS, IAM, CloudWatch, Lambda), Kubernetes, ArgoCD, Docker, Redis, "
            "Prometheus, Grafana, Pytest, Artifactory, GitHub."
        ),
    },
    {
        "from": "С 03.2021",
        "to": "До 12.2021",
        "role": "Инженер-программист",
        "project": (
            "Платформа интеграций доставки для ресторанов: единый внутренний API и кабинет поверх нескольких "
            "внешних площадок. Рестораны управляют меню, заказами, статусами, метриками и подключёнными "
            "провайдерами в одном месте. Нужны микросервисы, событийная модель и saga-оркестрация, чтобы "
            "заказы и статусы оставались согласованными при сбоях на стороне партнёров."
        ),
        "duties": [
            "Отвечал за интеграции с крупными площадками доставки (Uber Eats, Grubhub, DoorDash, GloriaFood).",
            "Свёл разные внешние API к единому внутреннему контракту: ретраи, rate limit, проверка исходящих credentials, устойчивость к смене поведения партнёров.",
            "Проектировал и развивал микросервисную архитектуру на GCP: Cloud Run, Pub/Sub, Datastore, Cloud Scheduler.",
            "Реализовал событийные потоки и saga-оркестрацию для согласованности заказов и статусов.",
            "Перевёл сервисы с Falcon/Starlette на FastAPI; REST-слой с OpenAPI; разграничил доступ кабинета ресторана к чужим данным площадки.",
            "Обеспечил ежедневные высоконагруженные синхронизации и сверку статусов под лимитами внешних API.",
            "Настроил кеширование в Redis, Docker, Sentry; документировал интеграционные контракты и регламенты поддержки.",
        ],
        "tech": (
            "Python, FastAPI, asyncio, GCP (Cloud Run, Pub/Sub, Datastore, Cloud Scheduler), Redis, Docker, "
            "OpenAPI/Swagger, Sentry, Pytest, Git."
        ),
    },
    {
        "from": "С 01.2020",
        "to": "До 02.2021",
        "role": "Инженер-программист",
        "project": (
            "Образовательная платформа-агрегатор услуг с контуром продаж, платежей и учёта. "
            "Клиенты публикуют учреждения, курсы и конкурсы; пользователи находят услугу и оформляют покупку. "
            "Внутри — витрина, биллинг, интеграция с бухгалтерией, согласованность проводок на денежных путях "
            "и миграция данных с legacy-контура без потери целостности."
        ),
        "duties": [
            "Разрабатывал backend на Django/DRF: каталог услуг, карточки клиентов, сценарии покупки и разграничение доступа к кабинетам.",
            "Спроектировал модель данных PostgreSQL под услуги, заказы, платежи и учётные статусы.",
            "Реализовал платёжный и биллинговый контур: статусы оплаты, защита от дублей и гонок на денежных путях.",
            "Обеспечил идемпотентность webhook/callback платёжных провайдеров и отклонение колбэков с невалидной подписью/контекстом.",
            "Интегрировал бухгалтерскую систему; обеспечил согласованность проводок и auditable-переходы финансовых сущностей.",
            "Написал и провёл миграционные скрипты переноса данных с контролем целостности и синхронизации.",
            "Реализовал фоновые задачи на Celery + Redis; кеш витрины в Redis; CI/CD, Docker Compose, Swagger, pytest, Sentry.",
        ],
        "tech": (
            "Python, Django, Django REST Framework, Celery, PostgreSQL, Redis, Docker, Docker Compose, "
            "Swagger API, GitLab CI/CD, Sentry, Pytest."
        ),
    },
]


def apply_content(doc: Document, *, rewrite_header: bool = True) -> None:
    if rewrite_header:
        set_runs_text(doc.paragraphs[0], "Максим В.")
        set_runs_text(doc.paragraphs[1], "Senior Python Backend Разработчик")
        ensure_spacer_after_title(doc)

    t0 = doc.tables[0]
    set_cell_lines(t0.rows[0].cells[0], SKILLS)
    set_cell_lines(t0.rows[0].cells[1], PROJECT_NAMES)
    set_cell_lines(
        t0.rows[1].cells[0],
        [
            "Образование",
            "Программное обеспечение информационных технологий, инженер-программист",
            "БГУИР, факультет компьютерных систем и сетей",
        ],
    )
    set_cell_lines(
        t0.rows[1].cells[1],
        ["Языковые навыки", "Русский, Английский (разговорный, B2+)"],
    )
    for row in t0.rows:
        for cell in row.cells:
            cleanup_empty_bullets(cell)

    set_cell_lines(doc.tables[1].rows[0].cells[0], SUMMARY)

    for idx, project in enumerate(PROJECTS):
        table = doc.tables[2 + idx]
        set_runs_text(table.rows[0].cells[1].paragraphs[0], project["from"])
        set_runs_text(table.rows[0].cells[2].paragraphs[0], project["to"])
        set_merged_value(table.rows[1], 1, project["role"])
        set_merged_value(table.rows[2], 1, project["project"])
        set_cell_lines(table.rows[3].cells[1], project["duties"])
        set_merged_value(table.rows[4], 1, project["tech"])

    body = doc.element.body
    while len(doc.tables) > 2 + len(PROJECTS):
        tbl = doc.tables[-1]._tbl
        body.remove(tbl)

    remove_square_before_projects_heading(doc)


def fill() -> Path:
    import shutil

    desktop_dir = Path.home() / "Desktop" / "CV" / "Вайтлейбл"
    desktop_dir.mkdir(parents=True, exist_ok=True)
    desktop_out = desktop_dir / "Василенко М. Senior CV.docx"

    doc = Document(str(TEMPLATE))
    apply_content(doc, rewrite_header=True)
    doc.save(str(OUT))
    doc.save(str(OUT_NAMED))
    shutil.copy2(OUT_NAMED, desktop_out)

    return OUT


if __name__ == "__main__":
    print(fill())
    print(OUT_NAMED)
    print(Path.home() / "Desktop" / "CV" / "Вайтлейбл" / "Василенко М. Senior CV.docx")
